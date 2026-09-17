from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).parent
SOURCE = ROOT / "Performance-Pay-Pay-Structure-Report.md"
OUTPUT = ROOT / "Performance-Pay-Pay-Structure-Report-Submission.docx"

NAVY = "1F4E79"
LIGHT_BLUE = "EAF2F8"
GRID = "D9D9D9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color=GRID):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relation_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)
    body = OxmlElement("w:t")
    body.text = text
    run.append(body)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_rich_text(paragraph, text):
    pattern = re.compile(r"\[([^\]]+)\]\((https?://[^)]+)\)")
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position:match.start()])
        add_hyperlink(paragraph, match.group(1), match.group(2))
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def prepare_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12

    for style_name, size in (("Title", 18), ("Heading 1", 14), ("Heading 2", 12)):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(14 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
    return doc


def table_from_lines(doc, lines):
    rows = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r"[-: ]+", cell) for cell in cells):
            continue
        rows.append(cells)
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.autofit = True
    repeat_table_header(table.rows[0])
    for row_index, row_values in enumerate(rows):
        for column_index, value in enumerate(row_values):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_border(cell)
            if row_index == 0:
                set_cell_shading(cell, NAVY)
            elif row_index % 2 == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            run.font.name = "Arial"
            run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
            run.font.size = Pt(10)
            if row_index == 0:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def main():
    text = SOURCE.read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = prepare_doc()
    index = 0
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        if line.startswith("# "):
            paragraph = doc.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run(line[2:])
        elif line.startswith("## "):
            doc.add_paragraph(line[3:], style="Heading 1")
        elif line.startswith("### "):
            doc.add_paragraph(line[4:], style="Heading 2")
        elif line.startswith("|"):
            end = index
            while end < len(lines) and lines[end].startswith("|"):
                end += 1
            table_from_lines(doc, lines[index:end])
            index = end
            continue
        elif line.startswith("**") and line.endswith("**"):
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(line[2:-2])
            run.bold = True
        elif line.startswith("**"):
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            content = line.replace("**", "")
            run = paragraph.add_run(content)
            run.italic = True
        else:
            paragraph = doc.add_paragraph()
            add_rich_text(paragraph, line)
        index += 1

    doc.add_paragraph("Sources", style="Heading 1")
    sources = [
        ("TCS Annual Report 2025 to 26", "https://www.tcs.com/content/dam/tcs/investor-relations/financial-statements/2025-26/ar/annual-report-2025-2026.pdf"),
        ("People Matters TCS variable pay report July 2025", "https://www.peoplematters.in/news/compensation-benefits/tcs-delivers-100percent-variable-pay-to-over-70percent-of-workforce-in-q1fy26-42114"),
        ("People Matters TCS pay structure report May 2026", "https://www.peoplematters.in/news/compensation-benefits/tcs-moves-part-of-employee-variable-pay-from-quarterly-to-annual-payouts-report-49899"),
        ("ICICI Bank AGM Notice FY2025 to 26", "https://www.sec.gov/Archives/edgar/data/1103838/000095010326011004/dp250272_6k.htm"),
        ("ICICI Bank Compensation Policy", "https://www.sec.gov/Archives/edgar/data/1103838/000095010326010820/dp249803_ex9701.htm"),
        ("Deloitte Benefits and Rewards", "https://www.deloitte.com/in/en/careers/deloitte-life/benefits.html"),
        ("About Deloitte USI", "https://www2.deloitte.com/ui/en/legal/about-deloitte.html"),
    ]
    for name, url in sources:
        paragraph = doc.add_paragraph(style="Normal")
        paragraph.paragraph_format.left_indent = Inches(0.2)
        paragraph.paragraph_format.first_line_indent = Inches(-0.2)
        add_hyperlink(paragraph, name, url)

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("In class Assignment 1")
    footer_run.font.name = "Arial"
    footer_run.font.size = Pt(9)
    doc.core_properties.title = "In class Assignment 1 Performance Pay Pay Structure and Salary Structure"
    doc.core_properties.author = ""
    doc.save(OUTPUT)


if __name__ == "__main__":
    main()
